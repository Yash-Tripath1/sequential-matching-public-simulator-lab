#!/usr/bin/env python3
"""Render research/RESEARCH_NOTE_DRAFT.md to DOCX using python-docx.

The source uses Markdown tables and image links plus horizontal rules as page breaks.
This is a lightweight project-specific renderer, not a general Markdown converter.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "research" / "RESEARCH_NOTE_DRAFT.md"
OUTPUT = ROOT / "research" / "RESEARCH_NOTE_DRAFT.docx"


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("__", "")
    text = text.replace("`", "").replace("*", "")
    return text.strip()


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def parse_table_row(line: str) -> list[str]:
    return [clean_inline(part) for part in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    parts = line.strip().strip("|").split("|")
    return bool(parts) and all(re.fullmatch(r"\s*:?-{3,}:?\s*", part) for part in parts)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def configure(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(0.58)
    sec.bottom_margin = Inches(0.58)
    sec.left_margin = Inches(0.68)
    sec.right_margin = Inches(0.68)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.1)
    normal.font.color.rgb = RGBColor(31, 48, 61)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.02

    for name, size, color in [("Title", 22, "142B3B"), ("Heading 1", 15, "142B3B"), ("Heading 2", 12, "24445C"), ("Heading 3", 10, "2C7A7B")]:
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.text = "SEQUENTIAL MATCHING  |  PUBLIC-SIMULATOR RESEARCH NOTE (DRAFT)"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(91, 111, 125)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Public synthetic results only  •  Page ")
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(91, 111, 125)
    add_page_field(footer)


def add_markdown_paragraph(doc: Document, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        return
    if stripped.startswith("# "):
        p = doc.add_paragraph(style="Title")
        p.add_run(clean_inline(stripped[2:]))
        return
    if stripped.startswith("## "):
        p = doc.add_paragraph(style="Heading 1")
        p.add_run(clean_inline(stripped[3:]))
        return
    if stripped.startswith("### "):
        p = doc.add_paragraph(style="Heading 2")
        p.add_run(clean_inline(stripped[4:]))
        return
    if stripped.startswith("#### "):
        p = doc.add_paragraph(style="Heading 3")
        p.add_run(clean_inline(stripped[5:]))
        return
    if stripped.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.right_indent = Inches(0.18)
        run = p.add_run(clean_inline(stripped[2:]))
        run.italic = True
        run.font.color.rgb = RGBColor(57, 81, 96)
        return
    if stripped.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(clean_inline(stripped[2:]))
        return
    if re.match(r"^\d+\.\s", stripped):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(clean_inline(re.sub(r"^\d+\.\s", "", stripped)))
        return
    if stripped.startswith("!["):
        match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if match:
            alt, rel_path = match.groups()
            img_path = (SOURCE.parent / rel_path).resolve()
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.add_run().add_picture(str(img_path), width=Inches(6.45))
            else:
                doc.add_paragraph(f"[Figure not found: {alt} — {rel_path}]")
        return
    p = doc.add_paragraph()
    p.add_run(clean_inline(stripped))
    if stripped.startswith("**Figure"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        for run in p.runs:
            run.italic = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(75, 94, 108)


def build_table(doc: Document, lines: list[str]) -> None:
    rows = [parse_table_row(line) for line in lines if not is_separator_row(line)]
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx in range(cols):
            value = row[cidx] if cidx < len(row) else ""
            cell = cells[cidx]
            cell.text = value
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(7.6)
                    if ridx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if ridx == 0:
                set_cell_shading(cell, "24445C")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main() -> None:
    doc = Document()
    configure(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip() == "---":
            if i != len(lines) - 1:
                doc.add_page_break()
            i += 1
            continue
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            build_table(doc, table_lines)
            continue
        add_markdown_paragraph(doc, line)
        i += 1
    doc.core_properties.title = "Budgeted Clarification for Sequential Matching — Public-Simulator Research Note (Draft)"
    doc.core_properties.subject = "Exploratory analysis of public synthetic simulator experiments"
    doc.core_properties.author = "Author(s) to be completed"
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
